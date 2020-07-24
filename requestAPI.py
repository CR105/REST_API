import json
import os
import requests
import zipfile
from collections import Counter
from datetime import datetime
from docx import Document
from docx.shared import Pt

slash = '/' #'\\'

def getToken(uri_Auth, user, password):
  headers_Auth = {
    'Content-Type': 'application/json',
    'Accept-API-Version': 'protocol=1.0, resource=2.0'
    }
  headers_Auth["X-OpenAM-Username"]= user
  headers_Auth["X-OpenAM-Password"]= password
  responseAuth = requests.request("POST", uri_Auth, headers=headers_Auth, data = {})
  if responseAuth.ok:
    responseAuth.close
    return responseAuth.json()["tokenId"]
  else:
    responseAuth.close
    return ''

def customAPIRequest(strNameEndPoint, data, token, methodReq, expectRC, path):
  request_dict = {}
  response_dict = {}
  evidence_dict = {}

  uri = data["host"] + data['endPoints'][strNameEndPoint]
  method = methodReq
  headers = data["headers"]
  headers["authorization"] = token
  headers["header.tracking_id"] = 'performance-' + strNameEndPoint + '-' + datetime.today().strftime('%Y%m%d')
  body = str(json.dumps(data["body"][strNameEndPoint], ensure_ascii=False))

  request_dict["headers"] = dict(headers)
  request_dict["body"] = json.loads(body)

  print("\nWaiting for...", strNameEndPoint.upper())
  response = requests.request(method, uri, headers=headers, data = body.encode('utf8'))

  if str(response.status_code) == expectRC:
    print(strNameEndPoint.upper(), ". Respuesta:",  str(response.status_code) + ' - ' + str(response.reason), "- PASSED", "- Elapsed:", str(response.elapsed))
    status = True
  else:
    print(strNameEndPoint.upper(), ". Respuesta:",  str(response.status_code) + ' - ' + str(response.reason), "- FEILED", "- Elapsed:", str(response.elapsed))
    status = False
  response_dict["statusCode"]= str(response.status_code) + ' - ' + str(response.reason)
  response_dict["headers"]= dict(response.headers)
  response_dict["body"] = response.json()
  response_dict["elapsed"] = str(response.elapsed)
  response.close()

  evidence_dict["request"] = dict(request_dict)
  evidence_dict["response"] = dict(response_dict)

  with open(path + slash + strNameEndPoint + "_" + str(response.status_code) +'.json', 'w', encoding='utf8') as json_file:
    json.dump(dict(evidence_dict), json_file, indent = 4, sort_keys=True)
    
  return status, response.status_code

def zipingResults(path, zipfilename):
  filesPath = []
  for root, directories, files in os.walk(path):
    for filename in files:
      filePath = os.path.join(root, filename)
      filesPath.append(filePath)
  
  zip_file = zipfile.ZipFile(path + slash + zipfilename, 'w')
  with zip_file:
    for filesName in filesPath:
      zip_file.write(filesName)

def report(data, event, endP, pathResource, pathExe, elapse):
  try:
    fileWord = open(pathResource + 'templateTestDocument.docx', 'rb')
    document = Document(fileWord)
    tablesInDoc = document.tables

    myStyle = document.styles['Normal']
    font = myStyle.font
    font.name = 'Calibri (Body)'
    font.size = Pt(10)

    exeResult = Counter(event)
    tablesInDoc[0].cell(0,1).text = data["report"]["title"]
    tablesInDoc[0].cell(1,1).text = data["report"]["author"]
    tablesInDoc[0].cell(2,1).text = data["report"]["phase"]
    tablesInDoc[0].cell(3,1).text = str(datetime.today().strftime('%Y/%m/%d %H:%M:%S'))
    tablesInDoc[0].cell(4,1).text = elapse
    

    tablesInDoc[1].cell(2,0).text = str(len(event))
    tablesInDoc[1].cell(2,1).text = str(exeResult[True] if True in exeResult else 0)
    tablesInDoc[1].cell(2,2).text = str(exeResult[False] if False in exeResult else 0)

    for idx in range(len(endP)):
        tablesInDoc[2].add_row()
        tablesInDoc[2].cell(idx + 1,0).text = str(idx + 1)
        tablesInDoc[2].cell(idx + 1,1).text = data["endPoints"][str(endP[idx][0])]
        tablesInDoc[2].cell(idx + 1,2).text = str(endP[idx][1])
        tablesInDoc[2].cell(idx + 1,3).text = 'PASSED' if event[idx] else 'FAILED'
    
    document.save(pathExe + slash +  str(data["report"]["title"]).replace(' ', '_') + '.docx')

  except (RuntimeError, TypeError, NameError):
    print("\nRuntime Error:", RuntimeError)
    print("\nType Error:", TypeError)
    print("\nName Error:", NameError)
  finally:
    fileWord.close()

def sendEmail():
  pass
